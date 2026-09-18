#!/usr/bin/env python
"""多格式数据接入 → 入库 → 检索 端到端测试（06）。

覆盖格式：
  文本类：md / txt / html / json
  表格类：csv / xlsx
  文档类：pdf / docx / pptx
  图片类：png（OCR）
  数据库：sqlite（表同步）

每种格式植入唯一标记 token，入库后检索该 token，校验能否命中正确的来源。
用法（在 backend 目录）：
  python scripts/06_multiformat_test.py
前置：API 运行中（默认 127.0.0.1:18000），Mongo/Redis/Kafka 可用。
"""

import json
import os
import sqlite3
import sys
import time
import urllib.error
import urllib.request
from pathlib import Path

BACKEND = Path(__file__).resolve().parent.parent
DATA = BACKEND / "data" / "e2e"
FILES = DATA / "files"
BASE = os.environ.get("API_BASE", "http://127.0.0.1:18000")

# 文件名 -> 唯一标记 token
CASES = [
    ("md", "e2e_sample.md", "E2E-MD-7788"),
    ("txt", "e2e_sample.txt", "E2E-TXT-9911"),
    ("html", "e2e_sample.html", "E2E-HTML-2233"),
    ("csv", "e2e_sample.csv", "E2E-CSV-4455"),
    ("xlsx", "e2e_sample.xlsx", "E2E-XLSX-6677"),
    ("json", "e2e_sample.json", "E2E-JSON-8899"),
    ("pdf", "e2e_sample.pdf", "E2E-PDF-1122"),
    ("docx", "e2e_sample.docx", "E2E-DOCX-3344"),
    ("pptx", "e2e_sample.pptx", "E2E-PPTX-5566"),
    ("png", "e2e_sample.png", "E2E-IMG-7788"),
]
DB_TOKEN = "E2E-DB-9900"
DB_NAME = "e2e.db"


# ---------------- 数据生成 ----------------

def gen_files() -> None:
    FILES.mkdir(parents=True, exist_ok=True)

    (FILES / "e2e_sample.md").write_text(
        "# E2E Markdown 文档\n\n唯一标记 E2E-MD-7788，用于验证 Markdown 解析与检索。\n\n"
        "## 章节\n本段包含中文内容，测试中文分词与向量召回。\n", encoding="utf-8")
    (FILES / "e2e_sample.txt").write_text(
        "E2E 纯文本文件\n唯一标记 E2E-TXT-9911，验证 txt 读取。\n", encoding="utf-8")
    (FILES / "e2e_sample.html").write_text(
        "<html><head><title>E2E 网页</title></head><body>"
        "<h1>E2E-HTML-2233</h1><p>这是 HTML 正文，验证标签剥离。</p></body></html>", encoding="utf-8")
    (FILES / "e2e_sample.csv").write_text(
        "id,name,note\n1,alpha,E2E-CSV-4455\n2,beta,second row\n", encoding="utf-8")
    (FILES / "e2e_sample.json").write_text(
        json.dumps([{"tag": "E2E-JSON-8899", "desc": "JSON 记录"}], ensure_ascii=False), encoding="utf-8")

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["id", "content"])
    ws.append([1, "E2E-XLSX-6677"])
    ws.append([2, "第二行数据"])
    wb.save(FILES / "e2e_sample.xlsx")

    import fitz  # PyMuPDF
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 100), "E2E PDF Document")
    page.insert_text((72, 130), "Unique token E2E-PDF-1122 for retrieval test.")
    doc.save(str(FILES / "e2e_sample.pdf"))
    doc.close()

    import docx
    d = docx.Document()
    d.add_heading("E2E Word 文档", level=1)
    d.add_paragraph("唯一标记 E2E-DOCX-3344，验证 docx 抽取。")
    d.save(FILES / "e2e_sample.docx")

    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    tb.text_frame.text = "E2E-PPTX-5566 幻灯片内容"
    prs.save(FILES / "e2e_sample.pptx")

    from PIL import Image, ImageDraw
    img = Image.new("RGB", (720, 160), "white")
    dr = ImageDraw.Draw(img)
    dr.text((20, 60), "E2E-IMG-7788 invoice total 4200", fill="black")
    img.save(FILES / "e2e_sample.png")

    # SQLite 演示库
    db = DATA / DB_NAME
    if db.exists():
        db.unlink()
    con = sqlite3.connect(db)
    con.execute("CREATE TABLE tickets(id INTEGER PRIMARY KEY, subject TEXT, amount REAL)")
    con.executemany("INSERT INTO tickets VALUES(?,?,?)", [
        (1, "E2E-DB-9900", 128.5), (2, "refund request", 60.0),
    ])
    con.commit()
    con.close()


# ---------------- API 封装 ----------------

def call(method, path, body=None, token=None, timeout=300):
    req = urllib.request.Request(BASE + path, method=method)
    req.add_header("Content-Type", "application/json")
    if token:
        req.add_header("Authorization", "Bearer " + token)
    data = json.dumps(body).encode() if body is not None else None
    try:
        with urllib.request.urlopen(req, data, timeout=timeout) as r:
            return r.status, json.loads(r.read().decode())
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode()


def login():
    s, r = call("POST", "/v1/auth/login", {"username": "admin", "password": "admin123"})
    return r["access_token"]


def ensure_datasource(token, name, type_, subtype, config):
    _, listing = call("GET", "/v1/datasources", token=token)
    for d in listing["datasources"]:
        if d["name"] == name:
            call("DELETE", "/v1/datasources/" + d["ds_id"], token=token)
    _, ds = call("POST", "/v1/datasources",
                 {"name": name, "type": type_, "subtype": subtype, "config": config,
                  "credentials": {}}, token)
    return ds["ds_id"]


def sync_and_wait(token, ds_id, mode="full", timeout=300):
    _, r = call("POST", f"/v1/datasources/{ds_id}/sync", {"mode": mode}, token)
    jid = r["job_id"]
    t0 = time.time()
    while time.time() - t0 < timeout:
        time.sleep(2)
        _, j = call("GET", f"/v1/jobs/{jid}", token=token)
        if j.get("status") in ("success", "failed"):
            return j
    return {"status": "timeout"}


def retrieve(token, query, top_n=5):
    _, r = call("POST", "/v1/retrieval/search",
                {"query": query, "top_k": 30, "top_n": top_n, "rerank": True}, token)
    return r.get("results", [])


# ---------------- 主流程 ----------------

def main() -> int:
    print("=" * 70)
    print("多格式入库 → 检索 端到端测试")
    print("=" * 70)
    gen_files()
    print(f"[生成] 测试文件目录: {FILES}")

    token = login()
    file_ds = ensure_datasource(token, "E2E-文件目录", "file", "directory", {"path": str(FILES)})
    db_ds = ensure_datasource(token, "E2E-SQLite", "database", "sqlite",
                              {"dialect": "sqlite", "path": str(DATA / DB_NAME)})
    print(f"[数据源] file={file_ds}  db={db_ds}")

    for ds_id, label in [(file_ds, "文件目录"), (db_ds, "SQLite")]:
        j = sync_and_wait(token, ds_id)
        print(f"[同步] {label}: status={j.get('status')} chunks={j.get('chunks')} failed={j.get('failed')}")

    # 检索校验
    results = []
    for fmt, fname, tok in CASES:
        hits = retrieve(token, tok)
        names = [h["doc_name"] for h in hits]
        ok = fname in names
        top = names[0] if names else "(无)"
        results.append((fmt, tok, fname, top, ok))
        print(f"[检索] {fmt:5s} {tok:14s} -> top={top:28s} {'✅' if ok else '❌ 期望 ' + fname}")

    # 数据库检索
    db_hits = retrieve(token, DB_TOKEN)
    db_names = [h["doc_name"] for h in db_hits]
    db_ok = "tickets" in db_names
    print(f"[检索] sqlite {DB_TOKEN:13s} -> top={db_names[0] if db_names else '(无)':28s} {'✅' if db_ok else '❌ 期望 tickets'}")

    # 语义检索（中文描述，非精确 token）
    sem = retrieve(token, "中文分词与向量召回")
    sem_ok = any("e2e_sample.md" in h["doc_name"] for h in sem[:3])
    print(f"[语义] 中文分词/向量召回 -> {[h['doc_name'] for h in sem[:3]]} {'✅' if sem_ok else '❌'}")

    # 图片 OCR 文本语义检索
    sem2 = retrieve(token, "invoice total 4200")
    sem2_ok = any("e2e_sample.png" in h["doc_name"] for h in sem2[:3])
    print(f"[语义] 图片 invoice 4200 -> {[h['doc_name'] for h in sem2[:3]]} {'✅' if sem2_ok else '❌'}")

    passed = (sum(1 for *_, ok in results if ok) + (1 if db_ok else 0)
              + (1 if sem_ok else 0) + (1 if sem2_ok else 0))
    total = len(results) + 3
    print("-" * 70)
    print(f"结果：{passed}/{total} 通过")
    return 0 if passed == total else 1


if __name__ == "__main__":
    sys.exit(main())
